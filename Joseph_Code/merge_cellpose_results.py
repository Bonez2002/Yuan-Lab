#!/usr/bin/env python3
"""Merge successful per-image Cellpose workbooks into one batch workbook."""

from __future__ import annotations

import argparse
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook


CORE_SHEETS = ("Object Results", "Image Summary", "Plasmid Metrics")
MAX_DATA_ROWS = 1_000_000


def experiment_from_workbook(path: Path, results_root: Path) -> str:
    relative = path.relative_to(results_root)
    return relative.parts[0]


def append_source_sheet(source, destination, experiment: str, expected_header):
    rows = source.iter_rows(values_only=True)
    header = tuple(next(rows))
    if expected_header is not None and header != expected_header:
        raise ValueError(
            f"Header mismatch in {source.title!r}: {header} != {expected_header}"
        )
    for row in rows:
        destination.append((experiment, *row))
    return header


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    success_markers = sorted(args.results_root.glob("*/per_image/*/_SUCCESS"))
    workbook_paths = [
        marker.parent / "3D_cell_intensity_results_CellposeSAM.xlsx"
        for marker in success_markers
    ]
    missing = [path for path in workbook_paths if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Successful tasks missing workbooks: {missing}")
    if not workbook_paths:
        raise ValueError(f"No successful per-image workbooks under {args.results_root}")

    output = Workbook(write_only=True)
    core_destinations = {
        name: output.create_sheet(name)
        for name in CORE_SHEETS
    }
    core_headers = {name: None for name in CORE_SHEETS}

    puncta_destinations = {}
    puncta_headers = {}
    puncta_counts = {}

    sources = output.create_sheet("Merge Sources")
    sources.append(("Experiment", "Workbook"))

    for index, path in enumerate(workbook_paths, start=1):
        experiment = experiment_from_workbook(path, args.results_root)
        print(f"[{index}/{len(workbook_paths)}] {experiment}: {path}")
        sources.append((experiment, str(path)))

        source_wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet_name in CORE_SHEETS:
                if sheet_name not in source_wb.sheetnames:
                    raise ValueError(f"{path} is missing worksheet {sheet_name!r}")
                source_ws = source_wb[sheet_name]
                rows = source_ws.iter_rows(values_only=True)
                header = tuple(next(rows))
                if core_headers[sheet_name] is None:
                    core_headers[sheet_name] = header
                    core_destinations[sheet_name].append(("Experiment", *header))
                elif header != core_headers[sheet_name]:
                    raise ValueError(f"Header mismatch in {path}: {sheet_name}")
                for row in rows:
                    core_destinations[sheet_name].append((experiment, *row))

            for sheet_name in source_wb.sheetnames:
                if not sheet_name.startswith("Puncta "):
                    continue

                source_ws = source_wb[sheet_name]
                rows = source_ws.iter_rows(values_only=True)
                header = tuple(next(rows))
                base = sheet_name[:31]

                if base not in puncta_headers:
                    puncta_headers[base] = header
                    puncta_counts[base] = 0
                    puncta_destinations[base] = []
                elif header != puncta_headers[base]:
                    raise ValueError(f"Puncta header mismatch in {path}: {sheet_name}")

                for row in rows:
                    part = puncta_counts[base] // MAX_DATA_ROWS + 1
                    while len(puncta_destinations[base]) < part:
                        suffix = "" if part == 1 else f" {part}"
                        title = (base[: 31 - len(suffix)] + suffix)[:31]
                        destination = output.create_sheet(title)
                        destination.append(("Experiment", *header))
                        puncta_destinations[base].append(destination)
                    puncta_destinations[base][part - 1].append((experiment, *row))
                    puncta_counts[base] += 1
        finally:
            source_wb.close()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    temporary = args.output.with_suffix(args.output.suffix + ".tmp")
    output.save(temporary)
    os.replace(temporary, args.output)
    print(f"Merged {len(workbook_paths)} workbooks into {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
