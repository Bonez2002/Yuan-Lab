"""
Cellpose-SAM / PyTorch GPU version of the paired 3D cell analysis workflow.

Group 1:
    segmented in 3D with Cellpose-SAM using PyTorch/CUDA.

Group 2:
    measured using the exact same Object_ID masks generated from Group 1.

Outputs:
    - One Excel workbook containing every object from every matched image pair
    - Raw Object_ID 3D TIFF
    - Colored label 3D TIFF
    - Colored overlay 3D TIFF

IMPORTANT:
    Edit only the USER SETTINGS section below.

Recommended environment:
    Python 3.10+
    NVIDIA GPU + a PyTorch build with CUDA support

Core packages:
    pip install cellpose tifffile numpy scipy openpyxl tqdm

Cellpose 4 uses PyTorch internally. The model selected below defaults to
"cpsam_v2", the current Cellpose-SAM model family.
"""

from __future__ import annotations

from pathlib import Path
import gc
import sys
import traceback
import inspect
import time
import platform
import atexit

import numpy as np
import tifffile
from scipy import ndimage as ndi
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from tqdm import tqdm

try:
    import torch
except Exception:
    torch = None

try:
    from cellpose import models
except Exception:
    models = None


TIFF_EXTENSIONS = {".tif", ".tiff"}



# ===================================================================
# WINDOWS ANTI-SLEEP
# ===================================================================

# Windows execution-state flags.
ES_CONTINUOUS = 0x80000000
ES_SYSTEM_REQUIRED = 0x00000001
ES_DISPLAY_REQUIRED = 0x00000002


def enable_anti_sleep(keep_display_on=False):
    """
    Prevent Windows from automatically sleeping while this Python process
    is running.

    keep_display_on=False:
        Keeps the computer/system awake, but still allows the monitor to turn
        off normally.

    keep_display_on=True:
        Keeps both the computer and display awake.

    On non-Windows systems this function safely does nothing.
    """
    if platform.system() != "Windows":
        print("Anti-sleep: non-Windows system detected; Windows anti-sleep skipped.")
        return False

    try:
        import ctypes

        flags = ES_CONTINUOUS | ES_SYSTEM_REQUIRED

        if keep_display_on:
            flags |= ES_DISPLAY_REQUIRED

        result = ctypes.windll.kernel32.SetThreadExecutionState(flags)

        if result == 0:
            raise OSError("SetThreadExecutionState returned 0.")

        if keep_display_on:
            print("Anti-sleep enabled: computer and display will stay awake.")
        else:
            print("Anti-sleep enabled: computer will stay awake; display may turn off.")

        return True

    except Exception as error:
        print(f"WARNING: could not enable Windows anti-sleep: {error}")
        return False


def disable_anti_sleep():
    """
    Restore normal Windows sleep behavior.
    """
    if platform.system() != "Windows":
        return

    try:
        import ctypes

        ctypes.windll.kernel32.SetThreadExecutionState(
            ES_CONTINUOUS
        )

        print("Anti-sleep disabled: normal Windows sleep behavior restored.")

    except Exception as error:
        print(f"WARNING: could not restore Windows sleep settings: {error}")

def get_settings():
    """
    ================================================================
    USER SETTINGS - CHANGE VALUES IN THIS SECTION ONLY
    ================================================================
    """
    class Settings:
        pass

    args = Settings()

    # ---------------------------------------------------------------
    # INPUT FOLDERS
    # ---------------------------------------------------------------

    # If this script is inside the experiment folder, leave this line alone.
    # args.root = Path(__file__).resolve().parent
    args.root = Path(r"S:/Lab Data/Stellaris 8/Plasmid Small Molecule/2026_08_12/Raw/Split")

    # ---------------------------------------------------------------
    # ANTI-SLEEP SETTINGS
    # ---------------------------------------------------------------

    # True = prevent Windows from automatically putting the computer to sleep
    # while this analysis script is running.
    args.prevent_sleep = True

    # False = computer stays awake, but monitor is still allowed to turn off.
    # True  = keep both computer AND monitor awake.
    #
    # False is recommended for long overnight analyses.
    args.keep_display_on = False

    args.group1_folder = "S:/Lab Data/Stellaris 8/Plasmid Small Molecule/2026_08_12/Raw/Split/group_1"
    args.group2_folder = "S:/Lab Data/Stellaris 8/Plasmid Small Molecule/2026_08_12/Raw/Split/group_2"

    # Expected matching filenames:
    # sample001_group1.tif
    # sample001_group2.tif
    args.group1_suffix = "_group1"
    args.group2_suffix = "_group2"

    # ---------------------------------------------------------------
    # OUTPUT LOCATION
    # ---------------------------------------------------------------

    # Change this to any folder you want.
    #
    # Windows example:
    # args.output_folder = Path(r"D:\Microscopy\Experiment_01\Cellpose_Results")
    #
    # Default:
    args.output_folder = args.root / "CellposeSAM_Results"

    args.output_filename = "3D_cell_intensity_results_CellposeSAM.xlsx"
    args.output = args.output_folder / args.output_filename

    # ---------------------------------------------------------------
    # CELLPOSE-SAM MODEL
    # ---------------------------------------------------------------

    # Current Cellpose-SAM model.
    args.pretrained_model = "cpsam_v2"

    # Use NVIDIA CUDA GPU when available.
    args.use_gpu = True

    # GPU index. Usually 0.
    args.gpu_device = 0

    # bfloat16 lowers GPU memory use on supported GPUs.
    args.use_bfloat16 = True

    # ---------------------------------------------------------------
    # TRUE 3D SEGMENTATION
    # ---------------------------------------------------------------

    # True = Cellpose performs 3D segmentation.
    args.do_3D = True

    # TIFFs in this workflow are expected as:
    # Z x Y x X
    args.z_axis = 0

    # None because Group 1 is a single-channel 3D stack.
    args.channel_axis = None

    # ---------------------------------------------------------------
    # VOXEL ANISOTROPY
    # ---------------------------------------------------------------

    # Set this if Z spacing differs from XY pixel spacing.
    #
    # anisotropy = Z voxel size / XY pixel size
    #
    # Example:
    # XY = 0.25 um/pixel
    # Z  = 1.00 um/plane
    # anisotropy = 1.00 / 0.25 = 4.0
    #
    # Leave None if voxels are close to isotropic or spacing is unknown.
    args.anisotropy = None

    # ---------------------------------------------------------------
    # CELL SIZE
    # ---------------------------------------------------------------

    # Cellpose-SAM is relatively robust to cell size.
    #
    # None = do not force a diameter.
    #
    # If cells are much larger than the model's typical scale, you can try
    # a value such as 40, 60, 80, etc.
    args.diameter = None

    # ---------------------------------------------------------------
    # SEGMENTATION SENSITIVITY
    # ---------------------------------------------------------------

    # Cell probability threshold.
    #
    # LOWER:
    #   finds more / larger masks
    #
    # HIGHER:
    #   is more conservative and can reduce weak false-positive regions
    #
    # Good starting point:
    args.cellprob_threshold = 0.0

    # Flow consistency threshold.
    #
    # This is primarily a mask-quality filter.
    args.flow_threshold = 0.4

    # Remove Cellpose objects smaller than this many voxels/pixels.
    args.min_size = 200

    # Maximum fraction of the entire image allowed for one mask.
    # Useful for rejecting catastrophic giant masks.
    args.max_size_fraction = 0.4

    # Number of Cellpose dynamics iterations.
    # None = automatic.
    # Increase for very large/long cells if needed.
    args.niter = None

    # ---------------------------------------------------------------
    # 3D FRAGMENTATION / OVERSEGMENTATION CONTROLS
    # ---------------------------------------------------------------

    # Smooth Cellpose flows in Z, Y, X before final 3D mask generation.
    #
    # This can help with fragmentation along Z.
    #
    # Examples:
    # args.flow3D_smooth = 0
    # args.flow3D_smooth = 1
    # args.flow3D_smooth = [2, 1, 1]   # stronger smoothing in Z
    args.flow3D_smooth = [2, 1, 1]

    # ---------------------------------------------------------------
    # INPUT NORMALIZATION
    # ---------------------------------------------------------------

    # Cellpose normalization.
    args.normalize = True

    # If cells are dark on a bright background, change to True.
    args.invert = False

    # ---------------------------------------------------------------
    # GPU BATCH / MEMORY
    # ---------------------------------------------------------------

    # Number of tiles evaluated at once.
    #
    # If GPU runs out of memory:
    #   reduce to 4, 2, or 1.
    #
    # If you have lots of VRAM:
    #   try 16 or 32.
    args.batch_size = 8

    # Tile overlap.
    args.tile_overlap = 0.1

    # For SAM models this should remain 256.
    args.bsize = 256

    # Test-time augmentation.
    # False is faster and recommended initially.
    args.augment = False

    # ---------------------------------------------------------------
    # OPTIONAL POST-PROCESSING
    # ---------------------------------------------------------------

    # Fill completely enclosed holes inside each final Cellpose object.
    args.fill_holes_after_cellpose = True

    # Maximum hole size to fill, in voxels.
    #
    # 0 means use complete binary hole filling inside each object.
    # Current implementation uses enclosed-hole filling.
    args.max_hole_size = 0

    # ---------------------------------------------------------------
    # NUCLEAR / CYTOSOL COMPARTMENTS
    # ---------------------------------------------------------------
    # Group 1 Cellpose objects are treated as nuclei.
    # True: expand every nucleus until nearest-neighbor territories meet.
    # This creates a full 3D Voronoi-like partition of the image.
    args.expand_cytosol_until_touching = True

    # IMPORTANT: change these to your microscope calibration.
    args.xy_pixel_size_um = 0.25
    args.z_spacing_um = 1.0

    # ---------------------------------------------------------------
    # GROUP 2 PUNCTA / HUB DETECTION
    # ---------------------------------------------------------------
    args.detect_group2_puncta = True

    # Difference-of-Gaussians spot sizes in microns.
    args.puncta_sigma_small_um = 0.35
    args.puncta_sigma_large_um = 0.80

    # Higher = fewer/brighter puncta; lower = more/dimmer puncta.
    args.puncta_threshold_std = 3.0

    # None = automatic robust threshold.
    args.puncta_fixed_threshold = None

    args.puncta_min_volume_um3 = 0.05
    args.puncta_max_volume_um3 = None

    # Split a hub at the nuclear boundary so nuclear and cytosolic portions
    # are measured separately.
    args.split_puncta_at_nuclear_boundary = True

    # Extra QC outputs.
    args.save_compartment_labels = True
    args.save_puncta_labels = True
    args.save_puncta_overlay = False

    # ---------------------------------------------------------------
    # OUTPUT IMAGES
    # ---------------------------------------------------------------

    # Raw integer Object_ID TIFF.
    # This is the authoritative label image.
    args.save_labels = False

    # Easy-to-see RGB label TIFF.
    args.save_colored_labels = True

    # Colored labels over original Group 1.
    args.save_overlay = False

    # Overlay opacity.
    args.overlay_alpha = 0.45

    return args


def check_dependencies(args):
    if torch is None:
        raise ImportError(
            "PyTorch is not installed. Install a PyTorch build appropriate "
            "for your computer/GPU before running this script."
        )

    if models is None:
        raise ImportError(
            "Cellpose is not installed. Install it with:\n"
            "pip install cellpose"
        )

    print(f"PyTorch version: {torch.__version__}")
    print(f"CUDA available to PyTorch: {torch.cuda.is_available()}")

    if args.use_gpu:
        if torch.cuda.is_available():
            gpu_index = int(args.gpu_device)
            if gpu_index >= torch.cuda.device_count():
                raise ValueError(
                    f"Requested GPU {gpu_index}, but PyTorch sees only "
                    f"{torch.cuda.device_count()} CUDA GPU(s)."
                )
            print(f"Using CUDA GPU {gpu_index}: {torch.cuda.get_device_name(gpu_index)}")
        else:
            print(
                "WARNING: use_gpu=True, but CUDA is not available to PyTorch.\n"
                "Cellpose will be initialized without CUDA acceleration."
            )


def normalize_pair_name(path: Path, suffix: str) -> str:
    stem = path.stem
    if not stem.lower().endswith(suffix.lower()):
        raise ValueError(
            f'Filename "{path.name}" does not end with expected suffix "{suffix}".'
        )
    return stem[:-len(suffix)]


def find_pairs(group1_dir, group2_dir, group1_suffix, group2_suffix):
    if not group1_dir.is_dir():
        raise FileNotFoundError(f"Group 1 folder not found: {group1_dir}")

    if not group2_dir.is_dir():
        raise FileNotFoundError(f"Group 2 folder not found: {group2_dir}")

    def collect(folder, suffix):
        result = {}

        for path in sorted(folder.iterdir()):
            if (
                path.is_file()
                and path.suffix.lower() in TIFF_EXTENSIONS
                and path.stem.lower().endswith(suffix.lower())
            ):
                pair_name = normalize_pair_name(path, suffix)

                if pair_name in result:
                    raise ValueError(
                        f'Duplicate pair name "{pair_name}" in {folder}'
                    )

                result[pair_name] = path

        return result

    group1 = collect(group1_dir, group1_suffix)
    group2 = collect(group2_dir, group2_suffix)

    if not group1:
        raise FileNotFoundError(
            f'No TIFF files ending with "{group1_suffix}" found in {group1_dir}'
        )

    missing_group2 = sorted(set(group1) - set(group2))
    missing_group1 = sorted(set(group2) - set(group1))

    if missing_group2 or missing_group1:
        messages = []

        if missing_group2:
            messages.append(
                "Missing Group 2 matches: " + ", ".join(missing_group2)
            )

        if missing_group1:
            messages.append(
                "Missing Group 1 matches: " + ", ".join(missing_group1)
            )

        raise ValueError("\n".join(messages))

    return [
        (name, group1[name], group2[name])
        for name in sorted(group1)
    ]


def open_stack(path):
    try:
        image = tifffile.memmap(path)
    except Exception:
        image = tifffile.imread(path)

    image = np.asarray(image)
    image = np.squeeze(image)

    if image.ndim != 3:
        raise ValueError(
            f"{path.name}: expected a 3D TIFF stack, got shape {image.shape}"
        )

    return image


def create_cellpose_model(args):
    """
    Create the Cellpose-SAM model once and reuse it for every image.

    Cellpose expects a torch.device object, not an integer GPU index.
    """

    print()
    print("Loading Cellpose-SAM model...")
    print(f"Requested model: {args.pretrained_model}")

    if args.use_gpu and torch.cuda.is_available():
        gpu_index = int(args.gpu_device)

        if gpu_index >= torch.cuda.device_count():
            raise ValueError(
                f"GPU {gpu_index} was requested, but PyTorch sees only "
                f"{torch.cuda.device_count()} CUDA GPU(s)."
            )

        device = torch.device(f"cuda:{gpu_index}")

        print(
            f"Using GPU {gpu_index}: "
            f"{torch.cuda.get_device_name(gpu_index)}"
        )
    else:
        device = torch.device("cpu")

        if args.use_gpu:
            print(
                "WARNING: GPU requested, but CUDA is unavailable. "
                "Using CPU."
            )
        else:
            print("Using CPU.")

    model_kwargs = {
        "gpu": bool(args.use_gpu and torch.cuda.is_available()),
        "pretrained_model": args.pretrained_model,
        "device": device,
        "use_bfloat16": bool(args.use_bfloat16),
    }

    # Filter constructor arguments too, because Cellpose releases can differ.
    constructor_parameters = inspect.signature(
        models.CellposeModel
    ).parameters

    filtered_model_kwargs = {
        key: value
        for key, value in model_kwargs.items()
        if key in constructor_parameters
    }

    ignored = sorted(
        set(model_kwargs) - set(filtered_model_kwargs)
    )

    if ignored:
        print(
            "Cellpose constructor does not support these settings "
            f"in this installation; ignoring: {', '.join(ignored)}"
        )

    model = models.CellposeModel(
        **filtered_model_kwargs
    )

    print(f"Cellpose device: {getattr(model, 'device', device)}")

    return model


def run_cellpose_3d(model, image, args):
    """
    Run Cellpose-SAM 3D segmentation.

    This function inspects the installed CellposeModel.eval() signature and
    sends only arguments supported by that exact Cellpose installation. This
    avoids errors such as:
        unexpected keyword argument 'invert'
    """

    tqdm.write("    Preparing Cellpose-SAM 3D segmentation...")

    eval_parameters = inspect.signature(
        model.eval
    ).parameters

    # Newer Cellpose versions may place invert inside the normalization
    # dictionary rather than accepting invert= directly.
    if "invert" in eval_parameters:
        normalization_setting = bool(args.normalize)
    else:
        normalization_setting = {
            "normalize": bool(args.normalize),
            "norm3D": True,
            "invert": bool(args.invert),
        }

    requested_kwargs = {
        "batch_size": int(args.batch_size),
        "channel_axis": args.channel_axis,
        "z_axis": args.z_axis,
        "normalize": normalization_setting,
        "invert": bool(args.invert),
        "diameter": args.diameter,
        "flow_threshold": float(args.flow_threshold),
        "cellprob_threshold": float(args.cellprob_threshold),
        "do_3D": bool(args.do_3D),
        "anisotropy": args.anisotropy,
        "min_size": int(args.min_size),
        "max_size_fraction": float(args.max_size_fraction),
        "niter": args.niter,
        "augment": bool(args.augment),
        "tile_overlap": float(args.tile_overlap),
        "bsize": int(args.bsize),
        "flow3D_smooth": args.flow3D_smooth,
        "progress": True,
    }

    eval_kwargs = {
        key: value
        for key, value in requested_kwargs.items()
        if key in eval_parameters
    }

    ignored = sorted(
        set(requested_kwargs) - set(eval_kwargs)
    )

    if ignored:
        tqdm.write(
            "    Installed Cellpose does not support these eval settings; "
            f"ignoring: {', '.join(ignored)}"
        )

    tqdm.write(
        "    Cellpose parameters being used: "
        + ", ".join(sorted(eval_kwargs))
    )

    start_time = time.perf_counter()

    result = model.eval(
        image,
        **eval_kwargs,
    )

    elapsed = time.perf_counter() - start_time

    tqdm.write(
        f"    Cellpose model.eval() finished in {elapsed:.1f} seconds."
    )

    # Different Cellpose releases can return different tuple lengths.
    # We only need the first return value: masks.
    if isinstance(result, tuple):
        masks = result[0]
    else:
        masks = result

    labels = np.asarray(
        masks,
        dtype=np.int32,
    )

    if labels.shape != image.shape:
        raise ValueError(
            f"Cellpose returned mask shape {labels.shape}, "
            f"but input image shape is {image.shape}."
        )

    return labels


def fill_holes_per_object(labels):
    """
    Fill enclosed holes while preserving Object_ID boundaries.

    Processing is done object-by-object within bounding boxes so that filling
    one cell does not merge it with a neighboring cell.
    """

    if int(labels.max()) == 0:
        return labels

    print("    Filling enclosed holes inside Cellpose objects...")

    output = labels.copy()
    boxes = ndi.find_objects(labels)

    for object_id, box in enumerate(boxes, start=1):
        if box is None:
            continue

        region = labels[box]
        object_mask = region == object_id

        filled = ndi.binary_fill_holes(object_mask)

        # Only add newly filled voxels where no other label currently exists.
        target = output[box]
        add = filled & (target == 0)
        target[add] = object_id

    return output


def relabel_consecutively(labels):
    unique = np.unique(labels)
    unique = unique[unique != 0]

    if unique.size == 0:
        return np.zeros(labels.shape, dtype=np.int32)

    lut = np.zeros(int(labels.max()) + 1, dtype=np.int32)
    lut[unique] = np.arange(
        1,
        unique.size + 1,
        dtype=np.int32,
    )

    return lut[labels]


def measure_objects(labels, group1, group2):
    """
    Measure voxel count and intensity sums one Z plane at a time to avoid
    unnecessary full-stack intensity copies.
    """

    n_objects = int(labels.max())

    counts = np.zeros(n_objects + 1, dtype=np.int64)
    group1_sum = np.zeros(n_objects + 1, dtype=np.float64)
    group2_sum = np.zeros(n_objects + 1, dtype=np.float64)

    for z in range(labels.shape[0]):
        lab = np.asarray(labels[z]).ravel()
        g1 = np.asarray(group1[z]).ravel()
        g2 = np.asarray(group2[z]).ravel()

        counts += np.bincount(
            lab,
            minlength=n_objects + 1,
        )

        group1_sum += np.bincount(
            lab,
            weights=g1,
            minlength=n_objects + 1,
        )

        group2_sum += np.bincount(
            lab,
            weights=g2,
            minlength=n_objects + 1,
        )

    return counts, group1_sum, group2_sum


def make_label_colors(labels):
    """
    Make deterministic RGB colors for visualization only.
    """

    n_objects = int(labels.max())

    palette = np.array([
        [230, 25, 75],
        [60, 180, 75],
        [255, 225, 25],
        [0, 130, 200],
        [245, 130, 48],
        [145, 30, 180],
        [70, 240, 240],
        [240, 50, 230],
        [210, 245, 60],
        [250, 190, 212],
        [0, 128, 128],
        [220, 190, 255],
        [170, 110, 40],
        [255, 250, 200],
        [128, 0, 0],
        [170, 255, 195],
        [128, 128, 0],
        [255, 215, 180],
        [0, 0, 128],
        [128, 128, 128],
    ], dtype=np.uint8)

    lookup = np.zeros(
        (n_objects + 1, 3),
        dtype=np.uint8,
    )

    if n_objects > 0:
        object_ids = np.arange(
            1,
            n_objects + 1,
        )

        lookup[1:] = palette[
            (object_ids - 1) % len(palette)
        ]

    return lookup[labels]


def normalize_group1_for_overlay(group1):
    """
    Percentile-normalize Group 1 to uint8 for visualization.
    """

    flat = group1.reshape(-1)
    step = max(
        1,
        flat.size // 1_000_000,
    )

    sample = np.asarray(
        flat[::step],
        dtype=np.float32,
    )

    low, high = np.percentile(
        sample,
        (1.0, 99.5),
    )

    if high <= low:
        low = float(np.min(sample))
        high = float(np.max(sample))

    if high <= low:
        return np.zeros(
            group1.shape,
            dtype=np.uint8,
        )

    output = np.empty(
        group1.shape,
        dtype=np.uint8,
    )

    scale = 255.0 / (high - low)

    for z in range(group1.shape[0]):
        plane = np.asarray(
            group1[z],
            dtype=np.float32,
        )

        plane = np.clip(
            (plane - low) * scale,
            0,
            255,
        )

        output[z] = plane.astype(
            np.uint8,
        )

    return output


def make_overlay(group1, colored_labels, labels, alpha):
    gray = normalize_group1_for_overlay(group1)

    overlay = np.repeat(
        gray[..., None],
        3,
        axis=-1,
    )

    del gray

    for z in range(labels.shape[0]):
        foreground = labels[z] > 0

        if not np.any(foreground):
            continue

        base = overlay[z]
        colors = colored_labels[z]

        blended = (
            (1.0 - alpha)
            * base[foreground].astype(np.float32)
            +
            alpha
            * colors[foreground].astype(np.float32)
        )

        base[foreground] = np.clip(
            blended,
            0,
            255,
        ).astype(np.uint8)

    return overlay



def voxel_volume_um3(args):
    return (
        float(args.xy_pixel_size_um)
        * float(args.xy_pixel_size_um)
        * float(args.z_spacing_um)
    )


def expand_nuclei(labels, args):
    """Expand nuclei until nearest-neighbor territories touch in 3D."""
    if int(labels.max()) == 0:
        empty = np.zeros(labels.shape, dtype=np.int32)
        return empty, empty.copy()

    print("    Expanding nuclei until neighboring territories touch...")
    background = labels == 0
    _, indices = ndi.distance_transform_edt(
        background,
        sampling=(
            float(args.z_spacing_um),
            float(args.xy_pixel_size_um),
            float(args.xy_pixel_size_um),
        ),
        return_indices=True,
    )
    # Each voxel inherits the Object_ID of its nearest nuclear voxel.
    whole = labels[indices[0], indices[1], indices[2]].astype(np.int32, copy=False)
    whole[labels > 0] = labels[labels > 0]
    cytosol = whole.copy()
    cytosol[labels > 0] = 0
    return whole, cytosol


def region_stats(region_labels, image):
    """Per-Object_ID intensity statistics."""
    n = int(region_labels.max())
    count = np.zeros(n + 1, dtype=np.int64)
    total = np.zeros(n + 1, dtype=np.float64)
    total2 = np.zeros(n + 1, dtype=np.float64)

    for z in range(region_labels.shape[0]):
        lab = region_labels[z].ravel()
        val = np.asarray(image[z], dtype=np.float64).ravel()
        count += np.bincount(lab, minlength=n + 1)
        total += np.bincount(lab, weights=val, minlength=n + 1)
        total2 += np.bincount(lab, weights=val * val, minlength=n + 1)

    mean = np.divide(
        total, count,
        out=np.full(n + 1, np.nan),
        where=count > 0,
    )
    var = np.divide(
        total2, count,
        out=np.full(n + 1, np.nan),
        where=count > 0,
    ) - mean * mean
    std = np.sqrt(np.maximum(var, 0))

    median = np.full(n + 1, np.nan)
    minimum = np.full(n + 1, np.nan)
    maximum = np.full(n + 1, np.nan)

    for oid, box in enumerate(ndi.find_objects(region_labels), start=1):
        if box is None:
            continue
        mask = region_labels[box] == oid
        if not np.any(mask):
            continue
        values = np.asarray(image[box])[mask]
        median[oid] = np.median(values)
        minimum[oid] = np.min(values)
        maximum[oid] = np.max(values)

    return {
        "count": count, "sum": total, "mean": mean, "median": median,
        "std": std, "min": minimum, "max": maximum,
    }


def detect_puncta(group2, whole_labels, nucleus_labels, args):
    """3D Difference-of-Gaussians Group 2 puncta detection."""
    z_um = float(args.z_spacing_um)
    xy_um = float(args.xy_pixel_size_um)
    s1 = float(args.puncta_sigma_small_um)
    s2 = float(args.puncta_sigma_large_um)

    sigma1 = (s1 / z_um, s1 / xy_um, s1 / xy_um)
    sigma2 = (s2 / z_um, s2 / xy_um, s2 / xy_um)

    print("    Enhancing Group 2 puncta with 3D Difference-of-Gaussians...")
    img = np.asarray(group2, dtype=np.float32)
    dog = (
        ndi.gaussian_filter(img, sigma1, mode="nearest")
        - ndi.gaussian_filter(img, sigma2, mode="nearest")
    )
    del img

    valid = whole_labels > 0
    values = dog[valid]
    step = max(1, values.size // 2_000_000)
    sample = values[::step]

    if args.puncta_fixed_threshold is None:
        med = float(np.median(sample))
        mad = float(np.median(np.abs(sample - med)))
        robust_sigma = 1.4826 * mad
        if robust_sigma <= 0:
            robust_sigma = float(np.std(sample))
        threshold = med + float(args.puncta_threshold_std) * robust_sigma
    else:
        threshold = float(args.puncta_fixed_threshold)

    print(f"    Puncta DoG threshold: {threshold:.6g}")
    candidate = (dog > threshold) & valid
    del dog, values, sample
    gc.collect()

    vv = voxel_volume_um3(args)
    min_vox = max(1, int(np.ceil(float(args.puncta_min_volume_um3) / vv)))
    max_vox = None
    if args.puncta_max_volume_um3 is not None:
        max_vox = int(np.floor(float(args.puncta_max_volume_um3) / vv))

    puncta_labels = np.zeros(candidate.shape, dtype=np.int32)
    records = []
    next_id = 1
    structure = ndi.generate_binary_structure(3, 1)

    for cell_id in np.unique(whole_labels[candidate]):
        if cell_id == 0:
            continue

        compartments = ("Nuclear", "Cytosolic")
        if not args.split_puncta_at_nuclear_boundary:
            compartments = ("WholeCell",)

        for compartment in compartments:
            mask = candidate & (whole_labels == cell_id)
            if compartment == "Nuclear":
                mask &= nucleus_labels == cell_id
            elif compartment == "Cytosolic":
                mask &= nucleus_labels == 0

            components, num = ndi.label(mask, structure=structure)
            sizes = np.bincount(components.ravel())

            for cid in range(1, num + 1):
                size = int(sizes[cid])
                if size < min_vox:
                    continue
                if max_vox is not None and size > max_vox:
                    continue

                pmask = components == cid
                puncta_labels[pmask] = next_id
                coords = np.argwhere(pmask)
                cz, cy, cx = coords.mean(axis=0)
                vals = np.asarray(group2)[pmask]

                records.append({
                    "Punctum_ID": next_id,
                    "Cell_Object_ID": int(cell_id),
                    "Compartment": compartment,
                    "Voxel_Count": size,
                    "Volume_um3": size * vv,
                    "Mean_Intensity": float(np.mean(vals)),
                    "Max_Intensity": float(np.max(vals)),
                    "Integrated_Intensity": float(np.sum(vals, dtype=np.float64)),
                    "Centroid_Z_vox": float(cz),
                    "Centroid_Y_vox": float(cy),
                    "Centroid_X_vox": float(cx),
                    "Centroid_Z_um": float(cz * z_um),
                    "Centroid_Y_um": float(cy * xy_um),
                    "Centroid_X_um": float(cx * xy_um),
                })
                next_id += 1

            del components

    # Signed centroid distance to the matching nuclear surface.
    boxes = ndi.find_objects(nucleus_labels)
    sampling = (z_um, xy_um, xy_um)
    # Territories can extend to the image edge, so allow full-image padding.
    pad_z = nucleus_labels.shape[0]
    pad_xy = max(nucleus_labels.shape[1], nucleus_labels.shape[2])

    for r in records:
        oid = r["Cell_Object_ID"]
        box = boxes[oid - 1] if oid - 1 < len(boxes) else None
        if box is None:
            r["Distance_From_Nuclear_Surface_um"] = np.nan
            continue

        z0 = max(0, box[0].start - pad_z)
        z1 = min(nucleus_labels.shape[0], box[0].stop + pad_z)
        y0 = max(0, box[1].start - pad_xy)
        y1 = min(nucleus_labels.shape[1], box[1].stop + pad_xy)
        x0 = max(0, box[2].start - pad_xy)
        x1 = min(nucleus_labels.shape[2], box[2].stop + pad_xy)

        nuc = nucleus_labels[z0:z1, y0:y1, x0:x1] == oid
        outside = ndi.distance_transform_edt(~nuc, sampling=sampling)
        inside = ndi.distance_transform_edt(nuc, sampling=sampling)

        z = int(np.clip(round(r["Centroid_Z_vox"]), z0, z1 - 1)) - z0
        y = int(np.clip(round(r["Centroid_Y_vox"]), y0, y1 - 1)) - y0
        x = int(np.clip(round(r["Centroid_X_vox"]), x0, x1 - 1)) - x0

        r["Distance_From_Nuclear_Surface_um"] = (
            -float(inside[z, y, x]) if nuc[z, y, x]
            else float(outside[z, y, x])
        )

    print(f"    Final Group 2 puncta/hubs: {len(records)}")
    return puncta_labels, records, threshold


def puncta_summary(records, n_objects):
    out = {}
    for oid in range(1, n_objects + 1):
        out[oid] = {}
        for c in ("Nuclear", "Cytosolic"):
            out[oid][c] = {
                "count": 0, "volume": 0.0, "total": 0.0,
                "means": [], "max": np.nan, "largest": 0.0,
            }

    for r in records:
        if r["Compartment"] not in ("Nuclear", "Cytosolic"):
            continue
        d = out[r["Cell_Object_ID"]][r["Compartment"]]
        d["count"] += 1
        d["volume"] += r["Volume_um3"]
        d["total"] += r["Integrated_Intensity"]
        d["means"].append(r["Mean_Intensity"])
        d["largest"] = max(d["largest"], r["Volume_um3"])
        d["max"] = (
            r["Max_Intensity"] if np.isnan(d["max"])
            else max(d["max"], r["Max_Intensity"])
        )
    return out


def make_compartment_rgb(nucleus_labels, cytosol_labels):
    rgb = np.zeros(nucleus_labels.shape + (3,), dtype=np.uint8)
    rgb[cytosol_labels > 0] = (0, 180, 220)
    rgb[nucleus_labels > 0] = (230, 60, 180)
    return rgb


def make_puncta_overlay(group2, nucleus_labels, cytosol_labels, puncta_labels):
    gray = normalize_group1_for_overlay(group2)
    rgb = np.repeat(gray[..., None], 3, axis=-1)
    for z in range(rgb.shape[0]):
        cyto = cytosol_labels[z] > 0
        nuc = nucleus_labels[z] > 0
        puncta = puncta_labels[z] > 0
        if np.any(cyto):
            base = rgb[z][cyto].astype(np.float32)
            rgb[z][cyto] = np.clip(
                0.8 * base + 0.2 * np.array([0, 180, 220]), 0, 255
            ).astype(np.uint8)
        if np.any(nuc):
            base = rgb[z][nuc].astype(np.float32)
            rgb[z][nuc] = np.clip(
                0.8 * base + 0.2 * np.array([230, 60, 180]), 0, 255
            ).astype(np.uint8)
        rgb[z][puncta] = (255, 255, 0)
    return rgb

def create_workbook():
    wb = Workbook()

    results = wb.active
    results.title = "Object Results"
    results.append([
        "Image", "Object_ID",
        "Nucleus_Voxel_Count", "Nucleus_Volume_um3",
        "Cytosol_Voxel_Count", "Cytosol_Volume_um3",
        "Whole_Cell_Voxel_Count", "Whole_Cell_Volume_um3",
        "Group1_Nuclear_Mean",
        "Group2_Nuclear_Mean", "Group2_Nuclear_Median",
        "Group2_Nuclear_StdDev", "Group2_Nuclear_Min",
        "Group2_Nuclear_Max", "Group2_Nuclear_Total",
        "Group2_Cytosol_Mean", "Group2_Cytosol_Median",
        "Group2_Cytosol_StdDev", "Group2_Cytosol_Min",
        "Group2_Cytosol_Max", "Group2_Cytosol_Total",
        "Group2_WholeCell_Mean",
        "Nuclear_to_Cytosol_Mean_Ratio",
        "Cytosol_to_Nuclear_Mean_Ratio",
        "Nuclear_Puncta_Count", "Nuclear_Puncta_Total_Volume_um3",
        "Nuclear_Puncta_Volume_Percent",
        "Nuclear_Puncta_Total_Intensity",
        "Nuclear_Puncta_Mean_Intensity",
        "Nuclear_Puncta_Max_Intensity",
        "Nuclear_Largest_Punctum_um3",
        "Cytosolic_Puncta_Count", "Cytosolic_Puncta_Total_Volume_um3",
        "Cytosolic_Puncta_Volume_Percent",
        "Cytosolic_Puncta_Total_Intensity",
        "Cytosolic_Puncta_Mean_Intensity",
        "Cytosolic_Puncta_Max_Intensity",
        "Cytosolic_Largest_Punctum_um3",
        "Z_Min", "Z_Max", "Y_Min", "Y_Max", "X_Min", "X_Max",
    ])

    puncta = wb.create_sheet("Puncta Results")
    puncta.append([
        "Image", "Cell_Object_ID", "Punctum_ID", "Compartment",
        "Voxel_Count", "Volume_um3", "Mean_Intensity", "Max_Intensity",
        "Integrated_Intensity", "Centroid_Z_vox", "Centroid_Y_vox",
        "Centroid_X_vox", "Centroid_Z_um", "Centroid_Y_um",
        "Centroid_X_um", "Distance_From_Nuclear_Surface_um",
    ])

    summary = wb.create_sheet("Image Summary")
    summary.append([
        "Image", "Group1_File", "Group2_File", "Cellpose_Model", "Device",
        "Object_Count", "Puncta_Count", "Puncta_DoG_Threshold",
        "Cytosol_Expansion_Method", "XY_Pixel_Size_um", "Z_Spacing_um", "Status",
    ])

    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)

    for ws in (results, puncta, summary):
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 18

    results.column_dimensions["A"].width = 42
    puncta.column_dimensions["A"].width = 42
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 48
    summary.column_dimensions["C"].width = 48
    return wb, results, puncta, summary


def main():
    args = get_settings()

    anti_sleep_enabled = False

    if args.prevent_sleep:
        anti_sleep_enabled = enable_anti_sleep(
            keep_display_on=bool(args.keep_display_on)
        )

        if anti_sleep_enabled:
            # Restore normal sleep behavior even if the program exits because
            # of an exception or Ctrl+C.
            atexit.register(disable_anti_sleep)

    check_dependencies(args)

    root = args.root.expanduser().resolve()

    group1_dir = root / args.group1_folder
    group2_dir = root / args.group2_folder

    output_path = Path(
        args.output
    ).expanduser().resolve()

    if int(args.min_size) < 1:
        raise ValueError(
            "min_size must be >= 1"
        )

    if float(args.xy_pixel_size_um) <= 0 or float(args.z_spacing_um) <= 0:
        raise ValueError("Physical voxel sizes must be > 0.")

    if float(args.puncta_sigma_small_um) <= 0:
        raise ValueError("puncta_sigma_small_um must be > 0.")

    if float(args.puncta_sigma_large_um) <= float(args.puncta_sigma_small_um):
        raise ValueError(
            "puncta_sigma_large_um must be greater than puncta_sigma_small_um."
        )

    if not 0.0 <= float(args.overlay_alpha) <= 1.0:
        raise ValueError(
            "overlay_alpha must be between 0 and 1"
        )

    pairs = find_pairs(
        group1_dir,
        group2_dir,
        args.group1_suffix,
        args.group2_suffix,
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    labels_dir = (
        output_path.parent
        / "labels"
    )

    if (
        args.save_labels
        or args.save_colored_labels
        or args.save_overlay
    ):
        labels_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

    print()
    print("=" * 72)
    print("CELLPOSE-SAM 3D ANALYSIS")
    print("=" * 72)
    print(f"Matched image pairs: {len(pairs)}")
    print(f"Model: {args.pretrained_model}")
    print(f"Output: {output_path}")

    # Load model ONCE and reuse it for every image.
    model = create_cellpose_model(args)

    workbook, results_ws, puncta_ws, summary_ws = create_workbook()

    total_objects = 0
    failed_pairs = 0

    overall_progress = tqdm(
        pairs,
        desc="Overall image progress",
        unit="image",
        dynamic_ncols=True,
    )

    for pair_index, (
        pair_name,
        group1_path,
        group2_path,
    ) in enumerate(
        overall_progress,
        start=1,
    ):

        overall_progress.set_postfix_str(
            pair_name[:35]
        )

        tqdm.write("")
        tqdm.write("=" * 72)
        tqdm.write(
            f"[{pair_index}/{len(pairs)}] {pair_name}"
        )

        try:
            group1 = open_stack(
                group1_path
            )

            group2 = open_stack(
                group2_path
            )

            if group1.shape != group2.shape:
                raise ValueError(
                    "Group 1 and Group 2 shapes do not match: "
                    f"{group1.shape} vs {group2.shape}"
                )

            print(
                f"  Shape Z,Y,X: {group1.shape}"
            )

            tqdm.write(
                f"  [{pair_index}/{len(pairs)}] "
                f"Starting Cellpose: {pair_name}"
            )

            labels = run_cellpose_3d(
                model,
                group1,
                args,
            )

            tqdm.write(
                f"  [{pair_index}/{len(pairs)}] "
                f"Cellpose finished: {pair_name}"
            )

            if args.fill_holes_after_cellpose:
                labels = fill_holes_per_object(
                    labels
                )

            labels = relabel_consecutively(
                labels
            )

            object_count = int(
                labels.max()
            )

            print(
                f"    Final objects: {object_count}"
            )

            if object_count == 0:
                summary_ws.append([
                    pair_name, group1_path.name, group2_path.name,
                    args.pretrained_model, str(model.device), 0, 0, None,
                    "Nearest nucleus until touching",
                    float(args.xy_pixel_size_um),
                    float(args.z_spacing_um),
                    "No objects detected",
                ])

                del group1
                del group2
                del labels

                gc.collect()

                if torch.cuda.is_available():
                    torch.cuda.empty_cache()

                continue

            print("    Expanding nuclei into simulated cytosol...")
            whole_labels, cytosol_labels = expand_nuclei(labels, args)

            print("    Measuring nuclear/cytosolic Group 2 signal...")
            nuc_g1 = region_stats(labels, group1)
            nuc_g2 = region_stats(labels, group2)
            cyto_g2 = region_stats(cytosol_labels, group2)
            whole_g2 = region_stats(whole_labels, group2)

            if args.detect_group2_puncta:
                puncta_labels, puncta_records, puncta_threshold = detect_puncta(
                    group2, whole_labels, labels, args
                )
            else:
                puncta_labels = np.zeros(labels.shape, dtype=np.int32)
                puncta_records = []
                puncta_threshold = np.nan

            ps = puncta_summary(puncta_records, object_count)
            vv = voxel_volume_um3(args)
            boxes = ndi.find_objects(labels)

            for object_id in range(1, object_count + 1):
                box = boxes[object_id - 1]
                if box is None:
                    continue

                z_slice, y_slice, x_slice = box
                nc = int(nuc_g2["count"][object_id])
                cc = int(cyto_g2["count"][object_id])
                wc = int(whole_g2["count"][object_id])
                nm = float(nuc_g2["mean"][object_id])
                cm = float(cyto_g2["mean"][object_id])

                ncr = nm / cm if np.isfinite(nm) and np.isfinite(cm) and cm != 0 else np.nan
                cnr = cm / nm if np.isfinite(nm) and np.isfinite(cm) and nm != 0 else np.nan

                npun = ps[object_id]["Nuclear"]
                cpun = ps[object_id]["Cytosolic"]
                nv = nc * vv
                cv = cc * vv
                npct = 100 * npun["volume"] / nv if nv > 0 else np.nan
                cpct = 100 * cpun["volume"] / cv if cv > 0 else np.nan

                results_ws.append([
                    pair_name, object_id,
                    nc, nv, cc, cv, wc, wc * vv,
                    float(nuc_g1["mean"][object_id]),
                    nm, float(nuc_g2["median"][object_id]),
                    float(nuc_g2["std"][object_id]),
                    float(nuc_g2["min"][object_id]),
                    float(nuc_g2["max"][object_id]),
                    float(nuc_g2["sum"][object_id]),
                    cm, float(cyto_g2["median"][object_id]),
                    float(cyto_g2["std"][object_id]),
                    float(cyto_g2["min"][object_id]),
                    float(cyto_g2["max"][object_id]),
                    float(cyto_g2["sum"][object_id]),
                    float(whole_g2["mean"][object_id]),
                    float(ncr), float(cnr),
                    npun["count"], npun["volume"], npct, npun["total"],
                    float(np.mean(npun["means"])) if npun["means"] else np.nan,
                    npun["max"], npun["largest"],
                    cpun["count"], cpun["volume"], cpct, cpun["total"],
                    float(np.mean(cpun["means"])) if cpun["means"] else np.nan,
                    cpun["max"], cpun["largest"],
                    z_slice.start, z_slice.stop - 1,
                    y_slice.start, y_slice.stop - 1,
                    x_slice.start, x_slice.stop - 1,
                ])

            for r in puncta_records:
                puncta_ws.append([
                    pair_name, r["Cell_Object_ID"], r["Punctum_ID"],
                    r["Compartment"], r["Voxel_Count"], r["Volume_um3"],
                    r["Mean_Intensity"], r["Max_Intensity"],
                    r["Integrated_Intensity"], r["Centroid_Z_vox"],
                    r["Centroid_Y_vox"], r["Centroid_X_vox"],
                    r["Centroid_Z_um"], r["Centroid_Y_um"],
                    r["Centroid_X_um"],
                    r.get("Distance_From_Nuclear_Surface_um", np.nan),
                ])

            # -------------------------------------------------------
            # SAVE RAW LABELS
            # -------------------------------------------------------

            if args.save_labels:
                label_path = (
                    labels_dir
                    / f"{pair_name}_labels.tif"
                )

                print(
                    f"    Saving: {label_path.name}"
                )

                tifffile.imwrite(
                    label_path,
                    labels.astype(
                        np.int32,
                        copy=False,
                    ),
                    photometric="minisblack",
                    metadata={
                        "axes": "ZYX"
                    },
                    compression="zlib",
                )

            # -------------------------------------------------------
            # SAVE COMPARTMENT / PUNCTA QC
            # -------------------------------------------------------
            if args.save_compartment_labels:
                comp_rgb = make_compartment_rgb(labels, cytosol_labels)
                tifffile.imwrite(
                    labels_dir / f"{pair_name}_NUCLEUS_CYTOSOL.tif",
                    comp_rgb, photometric="rgb",
                    metadata={"axes": "ZYXC"}, compression="zlib",
                )
                del comp_rgb

            if args.save_puncta_labels:
                tifffile.imwrite(
                    labels_dir / f"{pair_name}_GROUP2_PUNCTA.tif",
                    puncta_labels.astype(np.int32, copy=False),
                    photometric="minisblack",
                    metadata={"axes": "ZYX"}, compression="zlib",
                )

            if args.save_puncta_overlay:
                po = make_puncta_overlay(
                    group2, labels, cytosol_labels, puncta_labels
                )
                tifffile.imwrite(
                    labels_dir / f"{pair_name}_GROUP2_PUNCTA_OVERLAY.tif",
                    po, photometric="rgb",
                    metadata={"axes": "ZYXC"}, compression="zlib",
                )
                del po

            # -------------------------------------------------------
            # SAVE COLOR / OVERLAY
            # -------------------------------------------------------

            if (
                args.save_colored_labels
                or args.save_overlay
            ):
                colored_labels = make_label_colors(
                    labels
                )

                if args.save_colored_labels:
                    color_path = (
                        labels_dir
                        / f"{pair_name}_labels_COLOR.tif"
                    )

                    print(
                        f"    Saving: {color_path.name}"
                    )

                    tifffile.imwrite(
                        color_path,
                        colored_labels,
                        photometric="rgb",
                        metadata={
                            "axes": "ZYXC"
                        },
                        compression="zlib",
                    )

                if args.save_overlay:
                    overlay = make_overlay(
                        group1,
                        colored_labels,
                        labels,
                        float(args.overlay_alpha),
                    )

                    overlay_path = (
                        labels_dir
                        / f"{pair_name}_labels_OVERLAY.tif"
                    )

                    print(
                        f"    Saving: {overlay_path.name}"
                    )

                    tifffile.imwrite(
                        overlay_path,
                        overlay,
                        photometric="rgb",
                        metadata={
                            "axes": "ZYXC"
                        },
                        compression="zlib",
                    )

                    del overlay

                del colored_labels

            summary_ws.append([
                pair_name, group1_path.name, group2_path.name,
                args.pretrained_model, str(model.device), object_count,
                len(puncta_records),
                float(puncta_threshold) if np.isfinite(puncta_threshold) else None,
                "Nearest nucleus until touching",
                float(args.xy_pixel_size_um),
                float(args.z_spacing_um),
                "OK",
            ])

            total_objects += object_count

            print(
                f"  Completed: "
                f"{object_count} object(s)"
            )

            del group1, group2, labels, whole_labels, cytosol_labels
            del puncta_labels, puncta_records, ps
            del nuc_g1, nuc_g2, cyto_g2, whole_g2, boxes

            gc.collect()

            if torch.cuda.is_available():
                torch.cuda.empty_cache()

        except Exception as error:
            failed_pairs += 1

            print(
                f"  ERROR: {error}",
                file=sys.stderr,
            )

            traceback.print_exc()

            summary_ws.append([
                pair_name, group1_path.name, group2_path.name,
                args.pretrained_model, str(model.device), None, None, None,
                "Nearest nucleus until touching",
                float(args.xy_pixel_size_um),
                float(args.z_spacing_um),
                f"ERROR: {error}",
            ])

            gc.collect()

            if (
                torch is not None
                and torch.cuda.is_available()
            ):
                torch.cuda.empty_cache()

    # ---------------------------------------------------------------
    # EXCEL NUMBER FORMATTING
    # ---------------------------------------------------------------
    for ws in (results_ws, puncta_ws, summary_ws):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"

    print()
    print("Saving Excel workbook...")

    workbook.save(
        output_path
    )

    print()
    print("=" * 72)
    print("ANALYSIS COMPLETE")
    print("=" * 72)
    print(
        f"Matched pairs: {len(pairs)}"
    )
    print(
        f"Objects measured: {total_objects}"
    )
    print(
        f"Failed pairs: {failed_pairs}"
    )
    print(
        f"Excel file: {output_path}"
    )

    if (
        args.save_labels
        or args.save_colored_labels
        or args.save_overlay
    ):
        print(
            f"Label images: {labels_dir}"
        )

    return 1 if failed_pairs else 0


if __name__ == "__main__":
    try:
        raise SystemExit(
            main()
        )

    except KeyboardInterrupt:
        print(
            "\nCancelled by user.",
            file=sys.stderr,
        )
        raise SystemExit(130)

    except Exception as error:
        print(
            f"\nFATAL ERROR: {error}",
            file=sys.stderr,
        )
        raise SystemExit(1)